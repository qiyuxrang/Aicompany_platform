"""生成演示用工程清单 xlsx（列结构与平台契约一致）。"""
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

THIN = Side(style="thin", color="B7C2D0")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def build(target: Path, title: str, rows):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "工程量清单"

    sheet["A1"] = title
    sheet["A1"].font = Font(name="微软雅黑", size=14, bold=True)
    sheet.merge_cells("A1:E1")

    headers = ["序号", "项目名称", "单位", "数量", "备注"]
    for column, heading in enumerate(headers, start=1):
        cell = sheet.cell(row=2, column=column, value=heading)
        cell.font = Font(name="微软雅黑", size=10, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="2F5597")
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = BORDER

    for index, (name, unit, quantity, note) in enumerate(rows, start=1):
        values = [index, name, unit, quantity, note]
        for column, value in enumerate(values, start=1):
            cell = sheet.cell(row=2 + index, column=column, value=value)
            cell.font = Font(name="微软雅黑", size=10)
            cell.border = BORDER
            if column == 4:
                cell.number_format = "0.00"
            elif column != 2 and column != 5:
                cell.alignment = Alignment(horizontal="center")

    sheet.column_dimensions["A"].width = 8
    sheet.column_dimensions["B"].width = 34
    sheet.column_dimensions["C"].width = 10
    sheet.column_dimensions["D"].width = 12
    sheet.column_dimensions["E"].width = 20

    target.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(target)
    workbook.close()
    print(f"已生成 {target}（{len(rows)} 条）")


if __name__ == "__main__":
    output = Path(sys.argv[1])
    build(output / "榆林某园区弱电工程清单.xlsx", "榆林某园区弱电工程工程量清单（演示样张）", [
        ("配电箱 落地式 600×800×200", "台", 12, "含安装"),
        ("电力电缆 YJV-4×25+1×16", "m", 860, "铜芯"),
        ("槽式电缆桥架 200×100", "m", 420, "含支架"),
        ("网络摄像机 枪机 400万", "台", 48, "室外防水"),
        ("汇聚交换机 24口千兆", "台", 6, "含光模块"),
        ("机柜 42U 600×1000", "台", 3, "含PDU"),
    ])
    build(output / "榆林某园区给排水清单.xlsx", "榆林某园区给排水工程量清单（演示样张）", [
        ("镀锌钢管 DN100", "m", 320, "含管件"),
        ("闸阀 Z41H-16C DN100", "个", 18, "含法兰"),
        ("潜水排污泵 50WQ15-15", "台", 4, "一备一用"),
        ("混凝土 C30 现浇", "m3", 26, "基础"),
        ("土方开挖 三类土", "m3", 180, "含外运"),
    ])