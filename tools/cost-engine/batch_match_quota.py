# -*- coding: utf-8 -*-
"""批量套定额：清单逐条比对陕西省基价表，输出检索报告。

用法（在仓库根目录执行）：
    ./.venv/Scripts/python.exe tools/cost-engine/batch_match_quota.py

也可指定清单目录（默认取本目录 samples/）：
    ./.venv/Scripts/python.exe tools/cost-engine/batch_match_quota.py --dir <清单目录>

产出：
  1. 控制台分级报告（命中/待人工确认/无候选）
  2. 结果 xlsx，写到 output/ 下（该目录不入版本控制）

判定分级：
  - 命中候选   : score ≥ 0.70
  - 待人工确认 : 0.30 ≤ score < 0.70
  - 无候选     : score < 0.30 或无结果
注意：本表只给编号/名称/单位，**不含价格** —— D-05 未签认，价格不由本工具给出。
"""
import sys
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")

ROOT = Path(__file__).resolve().parent
sys.path.insert(0, str(ROOT))
import cost_cli  # noqa: E402

OUT_DIR = ROOT / "output"
OUT_XLSX = OUT_DIR / "定额套用结果.xlsx"

THIN = Side(style="thin", color="B7C2D0")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEAD_FILL = PatternFill("solid", fgColor="E8EEF6")
HIT_FILL = PatternFill("solid", fgColor="E6F2E6")
REVIEW_FILL = PatternFill("solid", fgColor="FDF3E0")
MISS_FILL = PatternFill("solid", fgColor="F6E7E7")

HIT, REVIEW = 0.70, 0.30


def grade(score):
    if score >= HIT:
        return "命中候选"
    if score >= REVIEW:
        return "待人工确认"
    return "无候选"


def top(entry_name, unit, n=3):
    """复用 cost_cli 的检索与排序规则，确保脚本与平台行为一致。"""
    return cost_cli.rank_candidates(entry_name, unit, limit=n)


def collect(rows_out, files):
    """读清单文件，逐条检索，返回结构化记录。"""
    records = []
    for path in files:
        filename = path.name
        wb = load_workbook(path, data_only=True)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        h_idx = next((i for i, r in enumerate(rows)
                      if r and any(c == "项目名称" for c in r if c)), None)
        if h_idx is None:
            print(f"  ⚠ 跳过 {filename}：未找到「项目名称」列，无法解析。")
            wb.close()
            continue
        header = [("" if c is None else str(c)) for c in rows[h_idx]]
        col = {t: i for i, t in enumerate(header)}
        # 列名可能是「数量」或「工程量」，两种都认
        qty_key = next((k for k in ("数量", "工程量") if k in col), None)
        for raw in rows[h_idx + 1:]:
            cells = [("" if c is None else str(c)) for c in raw]
            if not any(cells):
                continue
            item = cells[col["项目名称"]].strip()
            unit = cells[col["单位"]].strip() if "单位" in col else ""
            qty = cells[col[qty_key]].strip() if qty_key else ""
            if not item:
                continue
            cands = top(item, unit)
            best = cands[0][0] if cands else 0.0
            stem = cost_cli.normalize_query(item)
            records.append({
                "file": filename, "item": item, "unit": unit, "qty": qty,
                "stem": stem, "cands": cands, "best": best, "grade": grade(best),
                "remark": cells[col["备注"]].strip() if "备注" in col else "",
            })
            rows_out.append(records[-1])
        wb.close()
    return records


def report(records):
    print("=" * 100)
    print("定额套用报告 —— 清单逐条比对《陕西省建设工程基价表（2025）》")
    print("=" * 100)
    print(f"定额库规模：{len(cost_cli.load_quota_library())} 条")
    print("说明：仅给编号/名称/单位，不含价格（D-05 未签认，价格不由本工具给出）\n")

    stat = {"命中候选": 0, "待人工确认": 0, "无候选": 0}
    for r in records:
        stat[r["grade"]] += 1

    for r in records:
        print("-" * 100)
        print(f"清单条目：{r['item']}    （单位 {r['unit']}｜数量 {r['qty']}｜备注 {r['remark']}）")
        print(f"归一化后：{r['stem']}    判定：{r['grade']}（最高分 {r['best']:.2f}）")
        if not r["cands"]:
            print("   ✗ 无候选")
        for s, e in r["cands"]:
            print(f"   {s:>4.2f}  {e['code']:<11} {e['major'][:6]:<8} "
                  f"{e['name'][:56]:<58} {e['unit']}")

    print("\n" + "=" * 100)
    print("统计")
    print("=" * 100)
    total = len(records)
    for k in ["命中候选", "待人工确认", "无候选"]:
        print(f"  {k:<8} {stat[k]:>3} 条   {stat[k]/total*100:>5.1f}%")
    print(f"  {'合计':<8} {total:>3} 条")
    return stat


def write_xlsx(records):
    wb = Workbook()
    ws = wb.active
    ws.title = "定额套用结果"
    ws.append(["清单文件", "序号", "项目名称", "单位", "数量", "归一化查询词",
               "判定", "最高分", "候选编号", "候选名称（工序＋规格）", "候选单位",
               "候选专业", "单位是否一致"])
    for r in records:
        if r["cands"]:
            for i, (s, e) in enumerate(r["cands"]):
                ws.append([
                    r["file"], i + 1 if i == 0 else "", r["item"] if i == 0 else "",
                    r["unit"] if i == 0 else "", r["qty"] if i == 0 else "",
                    r["stem"] if i == 0 else "", r["grade"] if i == 0 else "",
                    round(s, 2), e["code"], e["name"], e["unit"], e["major"],
                    "是" if cost_cli.units_compatible(r["unit"], e["unit"]) else "否",
                ])
        else:
            ws.append([r["file"], 1, r["item"], r["unit"], r["qty"], r["stem"],
                       "无候选", "", "", "", "", "", ""])

    widths = [16, 5, 30, 6, 6, 20, 10, 7, 11, 52, 8, 16, 10]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[chr(64 + i) if i <= 26 else "A"].width = w
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = HEAD_FILL
        cell.border = BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center")
    fills = {"命中候选": HIT_FILL, "待人工确认": REVIEW_FILL, "无候选": MISS_FILL}
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.border = BORDER
            cell.alignment = Alignment(vertical="center", wrap_text=(cell.column == 10))
        g = row[6].value
        if g in fills:
            row[6].fill = fills[g]
    ws.freeze_panes = "A2"

    # 说明页
    ws2 = wb.create_sheet("口径说明")
    for line in [
        ["项目", "说明"],
        ["数据来源", "陕西省建设工程基价表（2025）—— 省住建厅公开文件转换稿"],
        ["条目规模", "29316 条，覆盖通用安装/市政/城市地下综合管廊/房屋建筑与装饰/园林绿化/绿色建筑"],
        ["检索方式", "清单条目归一化（剥离型号尺寸）后与定额条目名称比对，"
                     "含中文与字母的字命中率计分，单位一致加权 0.10"],
        ["判定分级", "命中候选 ≥0.70；待人工确认 0.30~0.70；无候选 <0.30"],
        ["重要：不含价格", "本表只给编号/名称/单位。D-05（定额数据授权、价格基准日、"
                           "费用税费口径）尚未签认，价格不由本工具输出。"],
        ["使用边界", "结果仅供人工筛选参考，不构成正式计价依据，不得直接用于投标或结算。"],
    ]:
        ws2.append(line)
    ws2.column_dimensions["A"].width = 18
    ws2.column_dimensions["B"].width = 96
    for cell in ws2[1]:
        cell.font = Font(bold=True)
        cell.fill = HEAD_FILL
    for row in ws2.iter_rows(min_row=1):
        for cell in row:
            cell.border = BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    wb.save(OUT_XLSX)
    return OUT_XLSX


def main(argv=None):
    argv = list(sys.argv[1:] if argv is None else argv)
    # --dir 指定清单所在目录；缺省用本目录 samples/
    if "--dir" in argv:
        at = argv.index("--dir")
        source = Path(argv[at + 1]).expanduser()
        if not source.is_dir():
            print(f"✗ 清单目录不存在：{source}")
            return 1
    else:
        source = ROOT / "samples"

    files = sorted(p for p in source.glob("*.xls*") if not p.name.startswith("~$"))
    if not files:
        print(f"✗ {source} 下未找到 xlsx/xls 清单。")
        return 1
    print(f"清单来源：{source}（{len(files)} 个文件）\n")

    records = []
    collect(records, files)
    if not records:
        print("✗ 未能从清单中解析出任何条目。")
        return 1

    report(records)
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    path = write_xlsx(records)
    print(f"\n已生成：{path}")
    print(f"（{path.stat().st_size:,} 字节）")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())