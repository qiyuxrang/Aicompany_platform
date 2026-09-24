from pathlib import Path
import csv
import hashlib
import json
import re
from datetime import datetime

from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT


root = Path(__file__).resolve().parents[1]
source = Path(r'C:\Users\BJRunner\Desktop\企业平台总体规划.docx')
files = [
    'CHANGELOG.md', 'specs/00-overview.md', 'specs/01-platform.md',
    'specs/02-product.md', 'specs/03-hr.md', 'specs/04-engineering.md',
    'specs/05-business.md', 'PLAN.md', 'TASKS.md', 'ACCEPTANCE.md', 'EXECUTION_STATUS.md', 'AI_HANDOFF.md',
]
texts = {name: (root / name).read_text(encoding='utf-8') for name in files}
requirements = {}
for name, text in texts.items():
    for match in re.finditer(r'^### ((?:PLT|PRD|HR|ENG|BIZ)-\d{3}) (.+)$', text, re.M):
        identifier, title = match.groups()
        if identifier in requirements:
            raise ValueError(f'Duplicate requirement: {identifier}')
        if f'AT-{identifier}' not in text:
            raise ValueError(f'Missing acceptance: {identifier}')
        requirements[identifier] = {'title': title, 'source': name}
task_rows = [line for line in texts['TASKS.md'].splitlines() if line.startswith('| T-')]
for identifier, item in requirements.items():
    item['tasks'] = [line.split('|')[1].strip() for line in task_rows if identifier in line]
    if not item['tasks']:
        raise ValueError(f'Missing task mapping: {identifier}')
    item['acceptance'] = f'AT-{identifier}'
    item['status'] = '本版未执行；历史实现与证据另见EXECUTION_STATUS.md'
references = set(re.findall(r'(?<![A-Z])(?:PLT|PRD|HR|ENG|BIZ)-\d{3}', '\n'.join(texts.values())))
if references - set(requirements):
    raise ValueError(f'Undefined requirements: {references - set(requirements)}')
if len(requirements) != 39:
    raise ValueError(f'Expected 39 requirements, got {len(requirements)}')
if len(task_rows) != 23:
    raise ValueError(f'Expected 23 tasks, got {len(task_rows)}')
old_root = root.parent / '企业平台SDD_20260921'
old_requirements = set()
for path in (old_root / 'specs').glob('*.md'):
    old_requirements.update(re.findall(r'^### ((?:PLT|PRD|HR|ENG|BIZ)-\d{3}) ', path.read_text(encoding='utf-8'), re.M))
if not old_requirements <= set(requirements):
    raise ValueError('A V0.1 requirement ID was removed')
protected = json.loads((root / 'qa/v01-source-manifest.json').read_text(encoding='utf-8-sig'))
for item in protected:
    if hashlib.sha256((old_root / item['file']).read_bytes()).hexdigest() != item['sha256']:
        raise ValueError(f"Protected source changed: {item['file']}")
expected_hash = '046907eb1ec6518c2e2d6f2a6ff0e6398c8e2ddbdf4c8903382bc6dc505860dd'
if hashlib.sha256(source.read_bytes()).hexdigest() != expected_hash:
    raise ValueError('Original document changed')

register = root / 'qa/acceptance-register.csv'
if register.exists():
    with register.open(encoding='utf-8-sig', newline='') as stream:
        existing = list(csv.DictReader(stream))
    if any(item.get('结论') != '本版未执行' or item.get('实际结果') or item.get('证据路径') for item in existing):
        raise ValueError('Refusing to overwrite executed acceptance results')
with register.open('w', encoding='utf-8-sig', newline='') as stream:
    writer = csv.writer(stream)
    writer.writerow(['需求ID', '验收ID', '标题', '规范文件', '任务', '结论', '负责人', '代码与环境版本', '实际结果', '证据路径'])
    for identifier, item in requirements.items():
        writer.writerow([identifier, item['acceptance'], item['title'], item['source'], ';'.join(item['tasks']), '本版未执行', '', '', '', ''])
(root / 'qa/traceability.json').write_text(json.dumps(requirements, ensure_ascii=False, indent=2), encoding='utf-8')

document = Document()
section = document.sections[0]
section.page_width, section.page_height = Cm(21), Cm(29.7)
section.top_margin, section.bottom_margin = Cm(1.8), Cm(1.8)
section.left_margin, section.right_margin = Cm(2), Cm(2)
section.header_distance, section.footer_distance = Cm(0.7), Cm(0.7)

def set_font(style, size, bold=False):
    style.font.name = 'Calibri'
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor.from_string('111111')
    style.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), 'Microsoft YaHei')

set_font(document.styles['Normal'], 10.5)
normal = document.styles['Normal'].paragraph_format
normal.space_after = Pt(5)
normal.line_spacing = 1.13
normal.widow_control = True
set_font(document.styles['Title'], 23, True)
set_font(document.styles['Subtitle'], 12)
document.styles['Subtitle'].font.italic = False
for style in document.styles:
    for border in style.element.xpath('.//w:pBdr'):
        border.getparent().remove(border)
set_font(document.styles['Heading 1'], 17, True)
set_font(document.styles['Heading 2'], 12, True)
set_font(document.styles['Heading 3'], 11, True)
for name in ['Heading 1', 'Heading 2', 'Heading 3']:
    paragraph_format = document.styles[name].paragraph_format
    paragraph_format.space_before = Pt(10)
    paragraph_format.space_after = Pt(6)
    paragraph_format.keep_with_next = True
    paragraph_format.keep_together = True
for name in ['List Bullet', 'List Number']:
    set_font(document.styles[name], 10.5)
    document.styles[name].paragraph_format.space_after = Pt(4)

header = section.header.paragraphs[0]
header.text = '企业 AI 业务协同平台  |  SDD V0.2 增量规范'
header.runs[0].font.size = Pt(8)
header.runs[0].font.color.rgb = RGBColor.from_string('666666')
footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run('V0.2  ·  2026-09-22  ·  第 ')
field = OxmlElement('w:fldSimple')
field.set(qn('w:instr'), 'PAGE')
footer._p.append(field)
footer.add_run(' 页')
for run in footer.runs:
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor.from_string('666666')

document.add_paragraph('企业 AI 业务协同平台实施规范', 'Title')
document.add_paragraph('业务任务型平台与可控成果交付', 'Subtitle')
document.add_paragraph('版本 0.2    修订日期 2026年9月22日    文档修订已获授权')
document.add_paragraph('保留现有平台与业务资产，不建设通用 WorkBuddy 替代品。以产品技术方案为首个闭环，深化验收条件、来源追溯、修改影响、版本留痕和质量回归；人事、经营及工程分批接入。')
document.add_paragraph('原37项需求保留，新增2项，映射23项任务。本次只更新文档，不修改业务代码；已有隔离成果继续保留，本版新增与深化验收尚未执行。代码范围、真实数据处理及正式上线须分别确认。')
document.add_paragraph('使用方式：先读修订说明与执行状态，再批准一个增量任务批次。开发以配套Markdown为准，本册为同源阅读快照。')

def inline(paragraph, text):
    for piece in re.split(r'(\*\*[^*]+\*\*|`[^`]+`)', text):
        if not piece:
            continue
        if piece.startswith('**') and piece.endswith('**'):
            run = paragraph.add_run(piece[2:-2])
            run.bold = True
        elif piece.startswith('`') and piece.endswith('`'):
            run = paragraph.add_run(piece[1:-1])
            run.font.name = 'Consolas'
            run.font.size = Pt(9)
        else:
            paragraph.add_run(piece)

def make_table(lines):
    rows = []
    for line in lines:
        values = [value.strip() for value in line.strip().strip('|').split('|')]
        if all(re.fullmatch(r':?-{3,}:?', value.replace(' ', '')) for value in values):
            continue
        rows.append(values)
    if not rows:
        return
    count = len(rows[0])
    if any(len(row) != count for row in rows):
        raise ValueError('Malformed Markdown table')
    table = document.add_table(rows=0, cols=count)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = {2: [4.2, 12.8], 3: [3.1, 6.9, 7.0], 4: [2.0, 5.3, 4.8, 4.9]}.get(count, [17/count]*count)
    for column, width in zip(table.columns, widths):
        column.width = Cm(width)
    table_properties = table._tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for edge in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        border = OxmlElement(f'w:{edge}')
        border.set(qn('w:val'), 'single')
        border.set(qn('w:sz'), '4')
        border.set(qn('w:color'), 'D9E0E5')
        borders.append(border)
    table_properties.append(borders)
    for row_index, values in enumerate(rows):
        row = table.add_row()
        properties = row._tr.get_or_add_trPr()
        no_split = OxmlElement('w:cantSplit')
        properties.append(no_split)
        if row_index == 0:
            repeat = OxmlElement('w:tblHeader')
            properties.append(repeat)
        for cell, value, width in zip(row.cells, values, widths):
            cell.width = Cm(width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(4)
            paragraph.paragraph_format.space_before = Pt(4)
            paragraph.paragraph_format.line_spacing = 1.08
            inline(paragraph, value)
            for run in paragraph.runs:
                run.font.size = Pt(9.5)
                if row_index == 0:
                    run.bold = True
            if row_index == 0:
                shade = OxmlElement('w:shd')
                shade.set(qn('w:fill'), 'EDF1F4')
                cell._tc.get_or_add_tcPr().append(shade)
    document.add_paragraph().paragraph_format.space_after = Pt(2)

for file_index, name in enumerate(files):
    lines = texts[name].splitlines()
    cursor = 0
    if name in ['specs/00-overview.md', 'PLAN.md', 'TASKS.md']:
        document.add_page_break()
    while cursor < len(lines):
        line = lines[cursor].strip()
        if not line:
            cursor += 1
            continue
        if line.startswith('|'):
            table_lines = []
            while cursor < len(lines) and lines[cursor].strip().startswith('|'):
                table_lines.append(lines[cursor])
                cursor += 1
            make_table(table_lines)
            continue
        heading = re.match(r'^(#{1,3}) (.+)', line)
        if heading:
            level = 1 if len(heading.group(1)) == 1 else 2
            title = heading.group(2)
            if level == 1:
                title = f'{file_index+1}  {title}'
            document.add_paragraph(title, f'Heading {level}')
        elif line.startswith('- '):
            inline(document.add_paragraph(style='List Bullet'), line[2:])
        elif re.match(r'^\d+\. ', line):
            inline(document.add_paragraph(), line)
        else:
            inline(document.add_paragraph(), line.removeprefix('> '))
        cursor += 1

properties = document.core_properties
properties.title = '企业 AI 业务协同平台实施规范 V0.2'
properties.subject = '总体规划评估、模块规范、实施任务及验收'
properties.author = '企业平台项目组'
properties.keywords = 'SDD,V0.2,业务任务,留痕,质量回归'
properties.comments = '由配套Markdown生成；历史成果保留，新版验收未执行。'
document.settings.element.append(OxmlElement('w:updateFields'))
document.settings.element[-1].set(qn('w:val'), 'true')
output = root / '企业平台总体规划_SDD_V0.2.docx'
document.save(output)
snapshot = {
    'generated_at': datetime.now().isoformat(timespec='seconds'),
    'source_document': str(source), 'source_sha256': expected_hash,
    'output_document': str(output),
    'requirement_count': len(requirements),
    'acceptance_count': len(requirements),
    'task_count': len(task_rows),
    'business_acceptance_status': 'V02_NOT_EXECUTED_HISTORY_PRESERVED',
    'version': '0.2', 'v01_requirement_ids_preserved': len(old_requirements),
    'v01_protected_files_unchanged': len(protected),
    'markdown_sha256': {name: hashlib.sha256((root/name).read_bytes()).hexdigest() for name in files},
    'docx_sha256': hashlib.sha256(output.read_bytes()).hexdigest(),
}
(root/'qa/build-report.json').write_text(json.dumps(snapshot, ensure_ascii=False, indent=2), encoding='utf-8')
print(json.dumps({'docx': str(output), 'requirements': len(requirements), 'tasks': len(task_rows)}, ensure_ascii=False))
