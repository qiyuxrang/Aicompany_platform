"""Fail-closed structural gate for platform-generated report deliverables.

This verifier reads final OOXML files, not generator-side counters.  It is used
by long-form acceptance runs and can also be run against files downloaded from
the product workbench.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import re
import sys
from pathlib import Path
from zipfile import BadZipFile, ZipFile

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "backend"))
from portal.product_formal import has_repeated_filler


W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
CHARACTER_COUNT_SCOPE = "chapter_body_non_whitespace_only_excludes_titles_diagrams_tables_appendices"
MERMAID_SOURCE = re.compile(rb"(?:flowchart|graph)\s+(?:TB|TD|BT|RL|LR)", re.I)
PROHIBITED_OUTPUT_COPY = (
    "待核草稿：未经正式内容、格式及产品负责人批准，不得作为正式方案使用。",
    "待核草稿 · 未获正式发布批准",
    "正式模板与样例、内容及格式批准尚未完成；本件始终为待核草稿。",
    "技术方案与可行性研究 · 待人工审核草稿",
    "内容来自平台当前报告版本，不继承正式业务批准",
    "平台自动生成的商务科技风待审草稿",
    "只使用当前报告来源；正式发布需要人工审核。",
)


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def chapter_body_characters(root, styles=None, *, paragraphs=None) -> int:
    body = root.find(W + "body")
    if body is None:
        return 0
    style_map = {} if styles is None else {
        style.get(W + "styleId"): style for style in styles.findall(W + "style")
    }
    started = False
    appendix = False
    count = 0
    for paragraph in body:
        if paragraph.tag != W + "p":
            continue
        text = "".join(node.text or "" for node in paragraph.findall(f"{W}r/{W}t"))
        level = paragraph.find(f"{W}pPr/{W}outlineLvl")
        style = paragraph.find(f"{W}pPr/{W}pStyle")
        style_id = style.get(W + "val", "") if style is not None else ""
        style_names = [style_id]
        visited = set()
        while style_id in style_map and style_id not in visited:
            visited.add(style_id)
            definition = style_map[style_id]
            name = definition.find(W + "name")
            style_names.append(name.get(W + "val", "") if name is not None else "")
            if level is None:
                level = definition.find(f"{W}pPr/{W}outlineLvl")
            parent = definition.find(W + "basedOn")
            style_id = parent.get(W + "val", "") if parent is not None else ""
        names = " ".join(style_names).lower()
        if re.search(r"toc|目录", names):
            continue
        heading = level is not None and level.get(W + "val") != "9"
        heading = heading or bool(re.search(r"heading|标题", names))
        if heading:
            started = True
            if re.search(r"(?:^|\s)(?:附录|附表|附件|appendix|appendices|annex)", text, re.I):
                appendix = True
            continue
        if not started or appendix:
            continue
        if re.search(r"caption|题注|图注|表注", names) or re.match(
            r"\s*(?:图\s*\d|表\s*\d|图说明|图示说明|图注|表注|附表|来源[：:]|注[：:]|figure\b|table\b)", text, re.I
        ):
            continue
        if any(paragraph.find(f".//{W}{tag}") is not None for tag in (
            "drawing", "pict", "object", "fldChar", "instrText", "txbxContent",
        )):
            continue
        visible = []
        for run in paragraph.findall(W + "r"):
            if any(run.find(f"{W}rPr/{W}{tag}") is not None for tag in ("vanish", "webHidden")):
                continue
            visible.append("".join(node.text or "" for node in run.findall(W + "t")))
        if paragraphs is not None:
            paragraphs.append("".join(visible))
        count += len(re.sub(r"\s+", "", "".join(visible)))
    return count


def inspect_docx(
    path: Path,
    *,
    minimum_characters: int,
    maximum_characters: int | None = None,
    minimum_figures: int,
    expected_header: str | None = None,
) -> dict:
    import xml.etree.ElementTree as ET

    issues: list[str] = []
    try:
        with ZipFile(path) as archive:
            names = archive.namelist()
            document_bytes = archive.read("word/document.xml")
            xml_parts = {name: archive.read(name) for name in names if name.endswith(".xml")}
            all_xml = b"".join(xml_parts.values())
            media = [name for name in names if name.startswith("word/media/")]
    except (OSError, BadZipFile, KeyError) as error:
        return {"path": str(path), "passed": False, "issues": [f"invalid_docx:{error}"]}

    root = ET.fromstring(document_bytes)
    text = "".join(node.text or "" for node in root.iter(W + "t"))
    styles = ET.fromstring(xml_parts["word/styles.xml"]) if "word/styles.xml" in xml_parts else None
    body_paragraphs = []
    characters = chapter_body_characters(root, styles, paragraphs=body_paragraphs)
    non_black_text_runs = []
    for name, payload in xml_parts.items():
        if not (name == "word/document.xml" or name.startswith("word/header")
                or name.startswith("word/footer")):
            continue
        part = ET.fromstring(payload)
        for run in part.iter(W + "r"):
            if not any((node.text or "").strip() for node in run.iter(W + "t")):
                continue
            color = run.find(f"{W}rPr/{W}color")
            if (color is None or (color.get(W + "val") or "").upper() != "000000"
                    or color.get(W + "themeColor") is not None):
                non_black_text_runs.append(name)
    outline = {0: 0, 1: 0}
    h1_page_breaks = 0
    body = root.find(W + "body")
    body_children = list(body) if body is not None else []
    body_positions = {id(node): index for index, node in enumerate(body_children)}
    first_h1_seen = False
    chapter_embedded_figures = 0
    for paragraph in root.iter(W + "p"):
        level = paragraph.find(f"{W}pPr/{W}outlineLvl")
        if first_h1_seen and paragraph.find(f".//{W}drawing") is not None:
            chapter_embedded_figures += 1
        if level is None:
            continue
        try:
            value = int(level.get(W + "val"))
        except (TypeError, ValueError):
            continue
        if value in outline:
            outline[value] += 1
        if value == 0:
            if first_h1_seen:
                position = body_positions.get(id(paragraph), -1)
                previous = body_children[position - 1] if position > 0 else None
                explicit_break = (
                    previous is not None and previous.tag == W + "p"
                    and any(node.get(W + "type") == "page"
                            for node in previous.iter(W + "br"))
                )
                style_break = paragraph.find(f"{W}pPr/{W}pageBreakBefore") is not None
                if explicit_break or style_break:
                    h1_page_breaks += 1
            first_h1_seen = True
    toc_depth_two = bool(re.search(rb'TOC[^<]{0,120}\\o\s+&quot;1-2&quot;', all_xml, re.I) or
                         re.search(rb'TOC[^<]{0,120}\\o\s+"1-2"', all_xml, re.I))
    if characters < minimum_characters:
        issues.append(f"characters_below_target:{characters}<{minimum_characters}")
    if maximum_characters is not None and characters > maximum_characters:
        issues.append(f"characters_above_target:{characters}>{maximum_characters}")
    if has_repeated_filler(body_paragraphs):
        issues.append("repeated_body_filler")
    length_issues = list(issues)
    if outline[0] < 1 or outline[1] < 1:
        issues.append(f"heading_depth_missing:h1={outline[0]},h2={outline[1]}")
    if not toc_depth_two:
        issues.append("toc_depth_two_missing")
    if non_black_text_runs:
        issues.append(f"non_black_text_runs:{len(non_black_text_runs)}")
    sections = list(root.iter(W + "sectPr"))
    for index, section in enumerate(sections, start=1):
        margins = section.find(W + "pgMar")
        expected = {"top": "1440", "bottom": "1440", "left": "1440", "right": "1440", "gutter": "0"}
        if margins is None or any(margins.get(W + key, "0") != value for key, value in expected.items()):
            issues.append(f"ordinary_margins_missing:section={index}")
    expected_chapter_breaks = max(0, outline[0] - 1)
    if h1_page_breaks != expected_chapter_breaks:
        issues.append(f"chapter_page_break_missing:{h1_page_breaks}/{expected_chapter_breaks}")
    headers = "".join(
        "".join(node.text or "" for node in ET.fromstring(payload).iter(W + "t"))
        for name, payload in xml_parts.items() if name.startswith("word/header")
    )
    expected_header = expected_header or ("技术方案" if minimum_figures == 15 else "可行性研究报告")
    if expected_header not in headers:
        issues.append(f"document_title_header_missing:{expected_header}")
    field_text = b" ".join(xml_parts.values()).decode("utf-8", errors="ignore")
    if not re.search(r"\bPAGE\b", field_text, re.I):
        issues.append("page_number_field_missing")
    for phrase in PROHIBITED_OUTPUT_COPY:
        if phrase in text or phrase in headers:
            issues.append(f"prohibited_output_copy:{phrase[:16]}")
    # The report cover contains one company logo in addition to report figures.
    if len(media) < minimum_figures + 1:
        issues.append(f"figures_below_target:{max(0, len(media)-1)}<{minimum_figures}")
    if chapter_embedded_figures < minimum_figures:
        issues.append(f"chapter_embedded_figures_below_target:{chapter_embedded_figures}<{minimum_figures}")
    if MERMAID_SOURCE.search(all_xml):
        issues.append("mermaid_source_leaked_into_docx")
    return {
        "path": str(path), "sha256": sha256(path), "bytes": path.stat().st_size,
        "characters_non_whitespace": characters, "heading1_count": outline[0],
        "character_count_scope": CHARACTER_COUNT_SCOPE,
        "minimum_characters": minimum_characters,
        "length_passed": not length_issues,
        "structural_passed": not issues[len(length_issues):],
        "heading2_count": outline[1], "toc_depth_two": toc_depth_two,
        "heading1_page_breaks": h1_page_breaks, "section_count": len(sections),
        "all_visible_word_text_black": not non_black_text_runs,
        "embedded_media": len(media), "minimum_figures": minimum_figures,
        "chapter_embedded_figures": chapter_embedded_figures,
        "mermaid_source_embedded": bool(MERMAID_SOURCE.search(all_xml)),
        "passed": not issues, "issues": issues,
    }


def inspect_pptx(path: Path) -> dict:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    issues: list[str] = []
    try:
        presentation = Presentation(path)
    except (OSError, BadZipFile, ValueError) as error:
        return {"path": str(path), "passed": False, "issues": [f"invalid_pptx:{error}"]}
    width, height = presentation.slide_width, presentation.slide_height
    slide_rows = []
    total_shapes = total_charts = total_connectors = total_pictures = 0
    source_marked = 0
    all_slide_text: list[str] = []
    for number, slide in enumerate(presentation.slides, start=1):
        texts: list[str] = []
        out_of_bounds: list[str] = []
        for shape in slide.shapes:
            total_shapes += 1
            if getattr(shape, "has_text_frame", False):
                texts.append(shape.text)
            if getattr(shape, "has_chart", False):
                total_charts += 1
            if shape.shape_type == MSO_SHAPE_TYPE.LINE:
                total_connectors += 1
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                total_pictures += 1
            if shape.left < 0 or shape.top < 0 or shape.left + shape.width > width or shape.top + shape.height > height:
                out_of_bounds.append(getattr(shape, "name", "shape"))
        combined = "\n".join(texts)
        all_slide_text.append(combined)
        if "内容来源：" in combined or "SHA256" in combined:
            source_marked += 1
        if not combined.strip() and not slide.shapes:
            issues.append(f"blank_slide:{number}")
        if out_of_bounds:
            issues.append(f"slide_out_of_bounds:{number}:{','.join(out_of_bounds[:5])}")
        slide_rows.append({"slide": number, "shapes": len(slide.shapes), "source_marked": bool("内容来源：" in combined or "SHA256" in combined)})
    slide_count = len(presentation.slides)
    if slide_count != 5:
        issues.append(f"slides_outside_fixed_target:{slide_count} != 5")
    if total_shapes < slide_count * 5:
        issues.append("insufficient_native_visual_structure")
    if total_connectors < 8:
        issues.append("native_connectors_missing")
    if total_charts < 1:
        issues.append("native_chart_missing")
    if total_pictures < 3:
        issues.append(f"insufficient_embedded_visual_assets:{total_pictures}<3")
    if source_marked < max(1, slide_count - 4):
        issues.append(f"source_marking_incomplete:{source_marked}/{slide_count}")
    if total_pictures >= total_shapes / 2:
        issues.append("deck_appears_image_flattened")
    joined_text = "\n".join(all_slide_text)
    for phrase in PROHIBITED_OUTPUT_COPY:
        if phrase in joined_text:
            issues.append(f"prohibited_copy_present:{phrase}")
    return {
        "path": str(path), "sha256": sha256(path), "bytes": path.stat().st_size,
        "slides": slide_count, "native_shapes": total_shapes, "native_charts": total_charts,
        "native_connectors": total_connectors, "pictures": total_pictures,
        "source_marked_slides": source_marked, "slide_inventory": slide_rows,
        "passed": not issues, "issues": issues,
    }


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--technical", required=True, type=Path)
    parser.add_argument("--feasibility", required=True, type=Path)
    parser.add_argument("--presentation", required=True, type=Path)
    parser.add_argument("--output", type=Path)
    args = parser.parse_args()
    result = {
        "schema": "PRODUCT_DELIVERABLE_QUALITY_V1",
        "acceptance_scope": "actual_formal_output_length_and_structure",
        "business_content_and_visual_review": "not_evaluated",
        "technical_solution": inspect_docx(
            args.technical, minimum_characters=50000,
            minimum_figures=15, expected_header="技术方案"),
        "feasibility": inspect_docx(
            args.feasibility, minimum_characters=70000,
            minimum_figures=20, expected_header="可行性研究报告"),
        "presentation": inspect_pptx(args.presentation),
    }
    result["passed"] = all(result[key]["passed"] for key in ("technical_solution", "feasibility", "presentation"))
    payload = json.dumps(result, ensure_ascii=False, indent=2)
    if args.output:
        args.output.parent.mkdir(parents=True, exist_ok=True)
        args.output.write_text(payload + "\n", encoding="utf-8")
    print(payload)
    return 0 if result["passed"] else 1


if __name__ == "__main__":
    sys.exit(main())
