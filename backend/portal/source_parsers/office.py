"""Native DOCX/XLSX/XLS extraction. No Office automation, formulas or links execute."""
import io
import re
from defusedxml import ElementTree as ET
from .core import LIMITS, ParseError, safe_zip, table_items

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
REL = "{http://schemas.openxmlformats.org/package/2006/relationships}"


def docx(content, result):
    with safe_zip(content, "word/document.xml") as archive:
        names = archive.namelist()
        for name in names:
            if name.endswith(".rels"):
                root = ET.fromstring(archive.read(name))
                if any(node.get("TargetMode") == "External" for node in root):
                    result.warn("external_link_ignored", "文档含外部链接，未访问或加载链接目标。", severity="info")
        parts = ["word/document.xml"] + sorted(name for name in names if re.fullmatch(r"word/(header\d+|footer\d+|footnotes|endnotes)\.xml", name))
        paragraph_number = table_number = 0
        def text(node):
            # Deleted revisions and field instructions are not final document text.
            if node.tag in {W + "del", W + "instrText"}: return ""
            if node.tag == W + "t": return node.text or ""
            if node.tag == W + "tab": return "\t"
            if node.tag in {W + "br", W + "cr"}: return "\n"
            return "".join(text(child) for child in node) + ('\n' if node.tag == W + 'p' else '')
        def walk(node, part):
            nonlocal paragraph_number, table_number
            if node.tag == W + "del": return
            if node.tag == W + "p":
                paragraph_number += 1
                result.add(text(node), {"part": part, "paragraph": paragraph_number})
            elif node.tag == W + "tbl":
                table_number += 1
                rows = []
                for row_index, row in enumerate(node.findall(W + "tr"), 1):
                    cells = [text(cell).strip() for cell in row.findall(W + "tc")]
                    if len(cells) > LIMITS["columns"] or row_index > LIMITS["rows"]:
                        raise ParseError("table_limit", "Word 表格超过行列上限，请拆分。")
                    block = result.add("\t".join(cells), {"part": part, "table": table_number, "row": row_index}, "table_row", cells=cells)
                    if block: rows.append(block)
                table_items(result, rows)
                if node.find(".//" + W + "vMerge") is not None or node.find(".//" + W + "gridSpan") is not None:
                    result.warn("merged_cells", "表格存在合并单元格，未向空白格自动填值，请核对行列对应。", {"part": part, "table": table_number})
            else:
                for child in node: walk(child, part)
        for part in parts:
            root = ET.fromstring(archive.read(part))
            if root.find(".//" + W + "ins") is not None or root.find(".//" + W + "del") is not None:
                result.warn("tracked_changes", "文档含修订，提取插入后的可见文字并排除删除内容，请核对最终版本。", {"part": part})
            walk(root, part)
        images = [name for name in names if name.startswith("word/media/") and name.lower().endswith((".png", ".jpg", ".jpeg", ".bmp", ".tif", ".tiff", ".webp"))]
        if images:
            from .visual import image
            for index, name in enumerate(images, 1):
                image(archive.read(name), result, {"part": "Word 内嵌图片", "image": index}, embedded=True)
        unsupported = [name for name in names if name.startswith("word/media/") and name not in images]
        if unsupported or any(name.startswith('word/charts/') for name in names): result.warn("unsupported_graphics", "部分内嵌矢量图、图表或媒体未转为文字，请查看原文件。", severity="partial")
        result.meta.update(paragraphs=paragraph_number, tables=table_number, images=len(images))


def xlsx(content, result):
    import openpyxl
    with safe_zip(content, "xl/workbook.xml") as archive:
        if any(name.startswith("xl/externalLinks/") for name in archive.namelist()):
            result.warn("external_link_ignored", "工作簿存在外部链接，未读取链接目标或刷新数值。", severity="info")
        merged = sum(archive.read(name).count(b"<mergeCell ") for name in archive.namelist() if name.startswith("xl/worksheets/") and name.endswith(".xml"))
        if merged: result.warn("merged_cells", "工作簿存在合并单元格，保持原位置，不向其他格自动填充值。")
        if any(name.startswith(('xl/media/', 'xl/charts/')) for name in archive.namelist()):
            result.warn('spreadsheet_graphics', '工作簿内嵌图片或图表未作语义解析，请将需要识别的图片单独上传。', severity='partial')
        for name in archive.namelist():
            if name.startswith('xl/worksheets/') and name.endswith('.xml'):
                root = ET.fromstring(archive.read(name))
                if any(node.get('hidden') in ('1', 'true') for node in root.iter()):
                    result.warn('hidden_rows_columns', '工作表包含隐藏行或列，本次按原单元格位置提取，使用前请核对隐藏内容。')
    workbook = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=False, keep_links=False)
    try:
        if len(workbook.worksheets) > 40: raise ParseError("sheet_limit", "工作表超过 40 个，请拆分工作簿。")
        count = 0
        for sheet in workbook.worksheets:
            if sheet.sheet_state != "visible":
                result.warn("hidden_sheet_skipped", "隐藏工作表未导入，请在原文件中取消隐藏后重新上传。", {"sheet": sheet.title}, "partial")
                continue
            if (sheet.max_row or 0) > LIMITS["rows"] or (sheet.max_column or 0) > LIMITS["columns"]:
                raise ParseError("table_limit", "工作表范围过大，请删除空白格式区域或拆分文件。")
            rows = []
            # Exporters sometimes under-report worksheet dimensions. Read actual
            # rows, enforcing limits again rather than silently dropping content.
            sheet.reset_dimensions()
            for row_index, row in enumerate(sheet.iter_rows(), 1):
                if row_index > LIMITS['rows'] or len(row) > LIMITS['columns']:
                    raise ParseError('table_limit', '工作表实际行列超过上限，请拆分文件。')
                if not row: continue
                count += len(row)
                if count > LIMITS["cells"]: raise ParseError("table_limit", "工作簿读取单元格数超过安全上限，请拆分。")
                cells, formulas = [], []
                for cell in row:
                    value = cell.value
                    if value is None: value = ""
                    elif hasattr(value, "isoformat"): value = value.isoformat()
                    elif isinstance(value, int) and not isinstance(value, bool) and re.fullmatch(r"0{2,20}", cell.number_format or ""):
                        value = str(value).zfill(len(cell.number_format))
                    else: value = str(value)
                    cells.append(value)
                    if cell.data_type == "f": formulas.append(cell.coordinate)
                    if cell.data_type == "e": result.warn("cell_error", "单元格包含 Excel 错误值，请修正原表。", {"sheet": sheet.title, "range": cell.coordinate})
                location = {"sheet": sheet.title, "row": row_index, "range": f"A{row_index}:{openpyxl.utils.get_column_letter(len(row))}{row_index}"}
                if formulas: result.warn("formula_unverified", "保留原公式，不执行计算或把缓存结果当作已核实数量：" + ", ".join(formulas), location)
                block = result.add("\t".join(cells), location, "table_row", cells=cells, formulas=formulas)
                if block: rows.append(block)
            table_items(result, rows)
        result.meta.update(sheets=len(workbook.worksheets), cells=count)
    finally: workbook.close()


def xls(content, result):
    import xlrd
    if not content.startswith(bytes.fromhex("d0cf11e0a1b11ae1")):
        raise ParseError("format_mismatch", "文件不是标准 XLS 工作簿，请另存为 XLSX。")
    try: book = xlrd.open_workbook(file_contents=content, on_demand=True)
    except Exception as error: raise ParseError("invalid_document", "XLS 已损坏、已加密或不是工作簿，请解密后另存为 XLSX。") from error
    try:
        count = 0
        if book.nsheets > 40: raise ParseError('sheet_limit', '工作表超过 40 个，请拆分工作簿。')
        for sheet in book.sheets():
            if sheet.visibility:
                result.warn("hidden_sheet_skipped", "隐藏工作表未导入。", {"sheet": sheet.name}, "partial")
                continue
            count += sheet.nrows * sheet.ncols
            if sheet.nrows > LIMITS["rows"] or sheet.ncols > LIMITS["columns"] or count > LIMITS["cells"]:
                raise ParseError("table_limit", "XLS 超过行列或单元格读取上限，请拆分。")
            rows = []
            for row in range(sheet.nrows):
                cells = []
                for col in range(sheet.ncols):
                    value = sheet.cell_value(row, col)
                    if isinstance(value, float) and value.is_integer(): value = int(value)
                    cells.append(str(value))
                block = result.add("\t".join(cells), {"sheet": sheet.name, "row": row + 1}, "table_row", cells=cells)
                if block: rows.append(block)
            # xlrd exposes cached values, not formula source. Keep as evidence only;
            # never auto-promote old XLS quantities to equipment facts.
        result.warn("legacy_xls_values", "旧 XLS 读取的是保存值，不能区分公式与常量；未自动写入设备数量。建议另存为 XLSX 后导入。")
        result.meta.update(sheets=book.nsheets, cells=count)
    finally: book.release_resources()
