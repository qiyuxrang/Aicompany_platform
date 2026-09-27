# /// script
# requires-python = ">=3.12,<3.13"
# dependencies = ["lxml==6.0.2", "python-docx==1.2.0", "python-pptx==1.0.2", "PyMuPDF==1.26.4", "pywin32==311; sys_platform == 'win32'"]
# ///
"""Installed Skill entry: typed JSON -> source-derived DOCX plus MD, HTML or PPTX."""
import argparse, json, re, html
from copy import deepcopy
from pathlib import Path
from lxml import etree as E
from ooxml import *
from contract import read_content
from prototypes import PrototypeLibrary

POLICY_PATH = ROOT/'assets'/'document-format-policy.json'

def ensure(parent, tag, first=False):
    node=parent.find(tag,NS)
    if node is None:
        node=E.Element(q(tag))
        parent.insert(0,node) if first else parent.append(node)
    return node

def set_on_off(parent, tag, value):
    for node in parent.findall(tag,NS): parent.remove(node)
    if value:
        E.SubElement(parent,q(tag)).set(q('w:val'),'1')

def format_paragraph(node, role, policy):
    """Apply only the user-approved interim typography; retain unspecified layout."""
    if node.tag!=q('w:p'): return node
    typography=policy['typography']
    spec=typography.get(role, typography['figure_and_table_caption']
                        if role=='table_caption' else typography['body'])
    pp=ensure(node,'w:pPr',first=True)
    spacing=ensure(pp,'w:spacing')
    if role=='figure':
        # An inline picture is part of the text line.  Applying the document's
        # exact 18 pt body grid clips a full-height diagram down to a thin band
        # in Microsoft Word.  Let this line expand to the picture's intrinsic
        # height, center it, and keep it with its following caption.
        # Set auto explicitly: the feasibility template's Normal style carries
        # an inherited exact line height, so merely deleting the direct value
        # still clips inline pictures.
        spacing.set(q('w:line'),'240'); spacing.set(q('w:lineRule'),'auto')
        spacing.set(q('w:before'),'120'); spacing.set(q('w:after'),'60')
        ensure(pp,'w:jc').set(q('w:val'),'center')
        set_on_off(pp,'w:keepNext',True); set_on_off(pp,'w:keepLines',True)
    else:
        spacing.set(q('w:line'),'360'); spacing.set(q('w:lineRule'),'exact')
    if role in ('heading1','heading2','heading3'):
        for item in pp.findall('w:numPr',NS)+pp.findall('w:ind',NS)+pp.findall('w:pStyle',NS): pp.remove(item)
        outline=ensure(pp,'w:outlineLvl'); outline.set(q('w:val'),str(int(role[-1])-1))
        jc=ensure(pp,'w:jc'); jc.set(q('w:val'),spec['alignment'])
    elif role in ('figure_and_table_caption','table_caption'):
        jc=ensure(pp,'w:jc'); jc.set(q('w:val'),'center')
        if role=='table_caption':
            # Keep table captions with their table.  Figure paragraphs already
            # keep the image with its caption; the caption must remain free to
            # break before the explanatory note to avoid nearly blank pages.
            set_on_off(pp,'w:keepNext',True); set_on_off(pp,'w:keepLines',True)
    size=str(round(spec['size_pt']*2)); bold=spec.get('bold')
    for run in node.findall('w:r',NS):
        rp=ensure(run,'w:rPr',first=True)
        fonts=ensure(rp,'w:rFonts')
        for key in ('ascii','hAnsi','eastAsia','cs'): fonts.set(q('w:'+key),spec['font'])
        ensure(rp,'w:sz').set(q('w:val'),size)
        ensure(rp,'w:szCs').set(q('w:val'),size)
        lang=ensure(rp,'w:lang'); lang.set(q('w:val'),'zh-CN'); lang.set(q('w:eastAsia'),'zh-CN')
        if bold is not None:
            set_on_off(rp,'w:b',bold); set_on_off(rp,'w:bCs',bold)
    return node

def normalize_page(section, policy):
    size=ensure(section,'w:pgSz')
    size.set(q('w:w'),'11906'); size.set(q('w:h'),'16838')
    margins=policy['page']['margins_mm']
    page_margins=ensure(section,'w:pgMar')
    for key in ('top','bottom','left','right','gutter'):
        page_margins.set(q('w:'+key),str(round(float(margins[key])*1440/25.4)))
    page_margins.set(q('w:header'),str(round(float(policy['page']['header_distance_mm'])*1440/25.4)))
    page_margins.set(q('w:footer'),str(round(float(policy['page']['footer_distance_mm'])*1440/25.4)))
    grid=ensure(section,'w:docGrid')
    grid.set(q('w:type'),'linesAndChars'); grid.set(q('w:linePitch'),'360')

def normalize_footer(story, policy):
    for p in story.findall('.//w:p',NS):
        if any(code.split()[0].upper()=='PAGE' for code in field_instructions(p)):
            ensure(p,'w:pPr',first=True)
            ensure(p.find('w:pPr',NS),'w:jc').set(q('w:val'),'center')
            format_paragraph(p,'header',policy)

def clear_section_headers(section):
    for node in section.findall('w:headerReference',NS)+section.findall('w:titlePg',NS): section.remove(node)


def apply_page_contract(doc, data, policy):
    """Apply the approved margins, title header and centered PAGE field."""
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Mm, Pt

    def clear(paragraph):
        for child in list(paragraph._p):
            if child.tag != qn('w:pPr'):
                paragraph._p.remove(child)

    def set_font(run, size):
        run.font.name='宋体'; run.font.size=Pt(size)
        props=run._element.get_or_add_rPr(); fonts=props.get_or_add_rFonts()
        for key in ('ascii','hAnsi','eastAsia','cs'): fonts.set(qn('w:'+key),'宋体')

    def page_field(paragraph):
        clear(paragraph); paragraph.alignment=WD_ALIGN_PARAGRAPH.CENTER
        run=paragraph.add_run(); set_font(run,9)
        for kind,text_value in (('begin',None),('instr',' PAGE '),('separate',None),('text','1'),('end',None)):
            if kind in ('begin','separate','end'):
                node=OxmlElement('w:fldChar'); node.set(qn('w:fldCharType'),kind)
            elif kind=='instr':
                node=OxmlElement('w:instrText'); node.set('{http://www.w3.org/XML/1998/namespace}space','preserve'); node.text=text_value
            else:
                node=OxmlElement('w:t'); node.text=text_value
            run._r.append(node)

    title=data['metadata']['title']
    margins=policy['page']['margins_mm']
    for index,section in enumerate(doc.sections):
        section.top_margin=Mm(margins['top']); section.bottom_margin=Mm(margins['bottom'])
        section.left_margin=Mm(margins['left']); section.right_margin=Mm(margins['right'])
        section.gutter=Mm(margins['gutter'])
        section.header_distance=Mm(policy['page']['header_distance_mm'])
        section.footer_distance=Mm(policy['page']['footer_distance_mm'])
        section.different_first_page_header_footer=(index==0)
        section.header.is_linked_to_previous=False
        section.footer.is_linked_to_previous=False
        header=section.header.paragraphs[0]
        clear(header); header.alignment=WD_ALIGN_PARAGRAPH.CENTER
        set_font(header.add_run(title),policy['typography']['header']['size_pt'])
        ppr=header._p.get_or_add_pPr(); borders=ppr.find(qn('w:pBdr'))
        if borders is None:
            borders=OxmlElement('w:pBdr'); ppr.append(borders)
        bottom=borders.find(qn('w:bottom'))
        if bottom is None:
            bottom=OxmlElement('w:bottom'); borders.append(bottom)
        bottom.set(qn('w:val'),'double'); bottom.set(qn('w:sz'),'8'); bottom.set(qn('w:space'),'1'); bottom.set(qn('w:color'),'000000')
        page_field(section.footer.paragraphs[0])
        if index==0:
            clear(section.first_page_header.paragraphs[0])
            clear(section.first_page_footer.paragraphs[0])


def fit_table_width(table, source_width, target_width):
    if source_width<=target_width: return
    if source_width>target_width*1.10: raise ValueError('table width exceeds A4 body area')
    ratio=target_width/source_width
    for node in table.findall('.//w:gridCol',NS)+table.findall('.//w:tcW',NS)+table.findall('.//w:tblW',NS):
        value=node.get(q('w:w'))
        if value and value.isdigit() and node.get(q('w:type'),'dxa')=='dxa':
            node.set(q('w:w'),str(max(1,round(int(value)*ratio))))

def set_toc_depth(root, level):
    """Let the actual heading hierarchy determine the generated TOC depth."""
    found=False; replacement=r'\o "1-'+str(level)+'"'
    for node in root.findall('.//w:instrText',NS):
        if node.text and re.search(r'\bTOC\b',node.text,re.I):
            found=True
            if re.search(r'\\o\s+"1-\d+"',node.text,re.I):
                node.text=re.sub(r'\\o\s+"1-\d+"',lambda _: replacement,node.text,flags=re.I)
            else:
                node.text=node.text.rstrip()+' '+replacement+' '
    for node in root.findall('.//w:fldSimple',NS):
        instruction=node.get(q('w:instr'),'')
        if re.search(r'\bTOC\b',instruction,re.I):
            found=True
            node.set(q('w:instr'),re.sub(r'\\o\s+"1-\d+"',lambda _: replacement,instruction,flags=re.I))
    if not found: raise ValueError('template TOC field not found')

def force_word_text_black(parts):
    """Force every visible WordprocessingML text run to explicit #000000."""
    theme_attributes=(q('w:themeColor'),q('w:themeTint'),q('w:themeShade'))
    following_color={q('w:spacing'),q('w:w'),q('w:kern'),q('w:position'),q('w:sz'),q('w:szCs'),
                     q('w:highlight'),q('w:u'),q('w:effect'),q('w:bdr'),q('w:shd'),q('w:fitText'),
                     q('w:vertAlign'),q('w:rtl'),q('w:cs'),q('w:em'),q('w:lang'),
                     q('w:eastAsianLayout'),q('w:specVanish'),q('w:oMath'),q('w:rPrChange')}
    for name,payload in list(parts.items()):
        if not name.startswith('word/') or not name.endswith('.xml'):
            continue
        root=xml(payload); changed=False
        # Runs without rPr must also be explicit so hyperlink/TOC character
        # styles cannot reintroduce theme blue in another viewer.
        for run in root.findall('.//w:r',NS):
            rp=run.find('w:rPr',NS)
            if rp is None:
                rp=E.Element(q('w:rPr')); run.insert(0,rp)
            color=rp.find('w:color',NS)
            if color is None:
                color=E.Element(q('w:color'))
                position=next((index for index,child in enumerate(rp) if child.tag in following_color),len(rp))
                rp.insert(position,color)
            color.set(q('w:val'),'000000')
            for attribute in theme_attributes:
                color.attrib.pop(attribute,None)
            changed=True
        if changed: parts[name]=dump(root)

def add_cover_identity(doc, data, policy):
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Mm, Pt
    identity=policy['company_identity']; logo=(ROOT/identity['logo_path']).resolve()
    if not logo.is_relative_to((ROOT/'assets').resolve()) or sha(logo)!=identity['logo_sha256']:
        raise ValueError('company logo integrity mismatch')
    title=next((p for p in doc.paragraphs if p.text.strip()==data['metadata']['title']),None)
    if title is None: raise ValueError('cover title paragraph not found')
    logo_paragraph=title.insert_paragraph_before(); logo_paragraph.alignment=WD_ALIGN_PARAGRAPH.CENTER
    logo_paragraph.paragraph_format.space_after=Pt(6)
    logo_paragraph.add_run().add_picture(str(logo),width=Mm(identity['cover_logo_width_mm']))

def bookmark_name(identifier):
    if re.fullmatch(r'[A-Za-z][A-Za-z0-9_]{0,39}',identifier): return identifier
    return 'BJ_'+hashlib.sha256(identifier.encode('utf-8')).hexdigest()[:32]

def chinese_number(value):
    if not 1<=value<10000: raise ValueError('Chinese heading number outside supported range')
    digits='零一二三四五六七八九'; units=['','十','百','千']; out=''; zero=False
    for index,digit in enumerate(str(value)):
        n=int(digit); position=len(str(value))-index-1
        if n:
            if zero: out+='零'
            out+=digits[n]+units[position]; zero=False
        elif out: zero=True
    return out[1:] if out.startswith('一十') else out

def trusted_pack(family):
    pack=ROOT/'assets'/family; profile=json.loads((pack/'profile.json').read_text(encoding='utf-8'))
    for name,key in [('template.docx','template_sha256'),('prototypes.xml','prototypes_sha256')]:
        if sha(pack/name)!=profile[key]: raise ValueError('pack integrity mismatch: '+name)
    return pack,profile

def word(data,out,content_dir):
    check_new_outputs(out, quality=True)
    with staged_outputs(out,Path(out).with_suffix('.quality.json')) as (staged,quality):
        report=_write_word(data,staged,content_dir)
    return report

def _write_word(data,out,content_dir):
    out=Path(out).resolve()
    if out.is_relative_to(ROOT/'assets'): raise ValueError('refusing to overwrite master assets')
    from contract import validate
    validate(data)
    pack,profile=trusted_pack(data['family']); policy=json.loads(POLICY_PATH.read_text(encoding='utf-8'))
    if data['family'] not in policy['applies_to']: raise ValueError('document format policy does not cover family')
    parts=load_package(pack/'template.docx'); audit_package(parts)
    root=xml(parts['word/document.xml']); body=root.find('w:body',NS); proto=PrototypeLibrary((pack/'prototypes.xml').read_bytes())
    marker=next(p for p in body.findall('w:p',NS) if text(p)=='{{body}}')
    body.remove(marker); insertion=len(body)-1
    metadata={'owner':'待确认',**data['metadata']}
    fixed_replace(root,metadata)
    heading_levels=[block['level'] for block in data['blocks'] if block['type']=='heading']
    set_toc_depth(root,max(heading_levels,default=1))
    for name in list(parts):
        if re.fullmatch(r'word/footer\d+\.xml',name):
            story=xml(parts[name]); fixed_replace(story,metadata); normalize_footer(story,policy)
            parts[name]=dump(story)
    def prototype(kind,key):
        return proto.get(kind,key)
    emitted=[]; figures=[]; bookmark_map={}; counters=[0,0,0]; table_counts={}; figure_counts={}; current_section=body.find('w:sectPr',NS)
    existing_sections=root.findall('.//w:sectPr',NS)
    for section in existing_sections:
        normalize_page(section,policy); clear_section_headers(section)
    front_break_needed=False
    def available_width():
        sz=current_section.find('w:pgSz',NS); mar=current_section.find('w:pgMar',NS)
        return int(sz.get(q('w:w')))-int(mar.get(q('w:left')))-int(mar.get(q('w:right')))-int(mar.get(q('w:gutter'),'0'))
    def ptype(key,txt):
        p=paragraph(prototype('paragraph',key),txt)
        return p
    def bookmark(p,ident):
        # Stable content IDs in document bookmarks; no user-provided OOXML.
        name=bookmark_name(ident); bookmark_map[ident]=name
        i=str(len(emitted)+1); start=E.Element(q('w:bookmarkStart')); start.set(q('w:id'),i); start.set(q('w:name'),name)
        end=E.Element(q('w:bookmarkEnd')); end.set(q('w:id'),i); p.insert(1 if p.find('w:pPr',NS) is not None else 0,start); p.append(end)
    if front_break_needed:
        # A one-section source would otherwise leak the body header onto cover/TOC pages.
        p=ptype('paragraph',''); pp=ensure(p,'w:pPr',first=True); front=deepcopy(current_section)
        for ref in front.findall('w:headerReference',NS): front.remove(ref)
        section_type=ensure(front,'w:type'); section_type.set(q('w:val'),'nextPage')
        pp.append(front); format_paragraph(p,'body',policy); emitted.append(p)
    first_body_heading=True
    for b in data['blocks']:
        typ=b['type']; nodes=[]
        if typ in ('heading','paragraph','runs','list'):
            key='heading'+str(b['level']) if typ=='heading' else ('list' if typ=='list' else 'paragraph')
            base=prototype('paragraph',key)
            if typ=='heading' and re.match(r'^(第[一二三四五六七八九十0-9]+[章节]|\d+(?:\.\d+){0,2}[.、 ]?)',b['text']):
                raise ValueError('heading numbering is generated; omit textual number')
            if typ=='runs':
                p=paragraph(base,'')
                for r in list(p.findall('w:r',NS)): p.remove(r)
                br=base.find('w:r/w:rPr',NS)
                for item in b['runs']:
                    r=E.SubElement(p,q('w:r')); rp=deepcopy(br) if br is not None else E.Element(q('w:rPr'))
                    for k,tag in [('bold','w:b'),('italic','w:i')]:
                        if k in item:
                            for x in rp.findall(tag,NS): rp.remove(x)
                            E.SubElement(rp,q(tag)).set(q('w:val'),'1' if item[k] else '0')
                    r.append(rp); append_run_text(r,item['text'])
                nodes=[p]
            elif typ=='list':
                auto=profile['paragraphs'][key]['automatic_numbering']
                nodes=[paragraph(base,('' if auto else f'（{i+1}）')+t) for i,t in enumerate(b['items'])]
            else:
                value=b['text']
                if typ=='heading':
                    level=b['level']; counters[level-1]+=1
                    for lower in range(level,3): counters[lower]=0
                    value='.'.join(str(n) for n in counters[:level])+' '+value
                nodes=[paragraph(base,value)]
        elif typ=='table':
            t=prototype('table',b['prototype']); tp=profile['tables'][b['prototype']]
            if not tp['variable_rows']: raise ValueError('complex merged form requires fixed topology adapter; expansion refused')
            if len(b['columns'])!=tp['columns']: raise ValueError('table columns differ from actual prototype')
            fit_table_width(t,tp['width_twips'],available_width())
            rows=t.findall('w:tr',NS); header=rows[0]; sample=rows[1] if len(rows)>1 else rows[0]; final_row=rows[-1]
            for row in rows: t.remove(row)
            for i,values in enumerate([b['columns']]+b['rows']):
                row=deepcopy(header if i==0 else final_row if i==len(b['rows']) else sample)
                pr=row.find('w:trPr',NS)
                if pr is None: pr=E.Element(q('w:trPr')); row.insert(0,pr)
                if i==0 and pr.find('w:tblHeader',NS) is None: E.SubElement(pr,q('w:tblHeader'))
                for cell,value in zip(row.findall('w:tc',NS),values):
                    cp=cell.find('w:p',NS)
                    for p in list(cell.findall('w:p',NS)): cell.remove(p)
                    cell.append(format_paragraph(paragraph(cp,value),'body',policy))
                t.append(row)
            units=[name+'：'+unit for name,unit in zip(b['columns'],b['units']) if unit not in ('不适用','无','—','-')]
            chapter=max(counters[0],1); table_counts[chapter]=table_counts.get(chapter,0)+1
            title=re.sub(r'^表\s*\d+(?:\.\d+)?\s*','',b['caption']).strip()
            caption=ptype('caption',f'表 {chapter}.{table_counts[chapter]} {title}'+('（单位：'+'；'.join(units)+'）' if units else '')); nodes=[caption,t]
        elif typ=='form':
            t=prototype('table',b['prototype']); rows=t.findall('w:tr',NS)
            for entry in b['cells']:
                if entry['row']>=len(rows): raise ValueError('form row outside source topology')
                cells=rows[entry['row']].findall('w:tc',NS)
                if entry['cell']>=len(cells): raise ValueError('form cell outside source topology')
                cell=cells[entry['cell']]; vm=cell.find('w:tcPr/w:vMerge',NS)
                if vm is not None and vm.get(q('w:val'),'continue')!='restart': raise ValueError('cannot write vertical-merge continuation cell')
                originals=cell.findall('w:p',NS)
                for p in originals: cell.remove(p)
                base=originals[0] if originals else E.Element(q('w:p'))
                for i,value in enumerate(entry['paragraphs']): cell.append(format_paragraph(paragraph(originals[min(i,len(originals)-1)] if originals else base,value),'body',policy))
            chapter=max(counters[0],1); table_counts[chapter]=table_counts.get(chapter,0)+1
            title=re.sub(r'^表\s*\d+(?:\.\d+)?\s*','',b['caption']).strip()
            nodes=[ptype('caption',f'表 {chapter}.{table_counts[chapter]} {title}'),t]
        elif typ=='figure':
            chapter=max(counters[0],1); figure_counts[chapter]=figure_counts.get(chapter,0)+1
            title=re.sub(r'^图\s*\d+(?:\.\d+)?\s*','',b['caption']).strip()
            p=ptype('figure',''); nodes=[p,ptype('caption',f'图 {chapter}.{figure_counts[chapter]} {title}')]
            path=(Path(content_dir)/b['path']).resolve()
            if not path.is_relative_to(Path(content_dir).resolve()) or not path.is_file(): raise ValueError('figure not found within content directory')
            if b['width_mm']*1440/25.4>available_width(): raise ValueError('figure overflow')
            figures.append((b['id'],path,b['width_mm'],b['alt']))
        elif typ=='section':
            next_section=prototype('section',b['prototype'])
            normalize_page(next_section,policy)
            # End the preceding section using its own geometry, then switch the terminal section.
            p=ptype('paragraph',''); pp=p.find('w:pPr',NS)
            if pp is None: pp=E.SubElement(p,q('w:pPr'))
            pp.append(deepcopy(current_section)); nodes=[p]
            body.remove(current_section); current_section=next_section; body.append(current_section)
        if typ=='heading' and b['level']==1:
            # Emit a literal page break.  Some lightweight DOCX previewers
            # ignore w:pageBreakBefore even though Word honours it; w:br is
            # widely supported and makes every chapter boundary visible.
            # The template already separates the front matter from chapter 1.
            # Add explicit breaks only between body chapters, otherwise a
            # refreshed multi-page TOC can leave a blank page before chapter 1.
            if not first_body_heading:
                page_break=E.Element(q('w:p')); break_pp=E.SubElement(page_break,q('w:pPr'))
                break_spacing=E.SubElement(break_pp,q('w:spacing'))
                for key,value in [('before','0'),('after','0'),('line','2'),('lineRule','exact')]:
                    break_spacing.set(q('w:'+key),value)
                break_run=E.SubElement(page_break,q('w:r')); break_rp=E.SubElement(break_run,q('w:rPr'))
                E.SubElement(break_rp,q('w:sz')).set(q('w:val'),'2')
                E.SubElement(break_run,q('w:br')).set(q('w:type'),'page')
                emitted.append(page_break)
            pp=nodes[0].find('w:pPr',NS)
            if pp is None: pp=E.Element(q('w:pPr')); nodes[0].insert(0,pp)
            # Avoid a duplicate blank page when the prototype also carries
            # pageBreakBefore; the explicit break above is the canonical one.
            for flag in pp.findall('w:pageBreakBefore',NS): pp.remove(flag)
            first_body_heading=False
        role='heading'+str(b['level']) if typ=='heading' else 'body'
        for index,node in enumerate(nodes):
            if node.tag==q('w:p'):
                caption=((typ in ('table','form') and index==0) or (typ=='figure' and index==1))
                node_role=('table_caption' if typ in ('table','form') and index==0
                           else 'figure_and_table_caption' if typ=='figure' and index==1
                           else 'figure' if typ=='figure' and index==0
                           else role)
                format_paragraph(node,node_role,policy)
        for p in nodes:
            if p.tag==q('w:p'): bookmark(p,b['id']); break
        emitted.extend(nodes)
    # Pending remains in source JSON; material limitations belong in authored prose.
    if emitted and emitted[-1].tag==q('w:tbl'):
        # Word requires a closing paragraph after a final table; an inherited
        # large empty paragraph otherwise creates a spurious blank last page.
        tail=E.Element(q('w:p')); pp=E.SubElement(tail,q('w:pPr'))
        sp=E.SubElement(pp,q('w:spacing'))
        for key,value in [('before','0'),('after','0'),('line','20'),('lineRule','exact')]: sp.set(q('w:'+key),value)
        rp=E.SubElement(pp,q('w:rPr')); E.SubElement(rp,q('w:sz')).set(q('w:val'),'2')
        emitted.append(tail)
    for i,node in enumerate(emitted): body.insert(insertion+i,node)
    settings=xml(parts['word/settings.xml'])
    update=ensure(settings,'w:updateFields')
    update.set(q('w:val'),'true')
    parts['word/settings.xml']=dump(settings)
    parts['word/document.xml']=dump(root); audit_package(parts); save_package(out,parts)
    if figures or policy.get('company_identity'):
        from docx import Document
        from docx.shared import Mm
        doc=Document(out)
        add_cover_identity(doc,data,policy)
        for ident,path,width,alt in figures:
            node=next(p for p in doc.element.body.findall(q('w:p')) if any(x.get(q('w:name'))==bookmark_name(ident) for x in p.findall(q('w:bookmarkStart'))))
            from docx.text.paragraph import Paragraph
            p=Paragraph(node,doc._body); shape=p.add_run().add_picture(str(path),width=Mm(width)); shape._inline.docPr.set('descr',alt)
        apply_page_contract(doc,data,policy)
        doc.save(out)
    final_parts=load_package(out); force_word_text_black(final_parts)
    audit_package(final_parts); save_package(out,final_parts)
    report={'structural_pass':True,'rendered':False,'visually_reviewed':False,'unverified_items':['Office pagination/TOC refresh and visual inspection required','Engineering assertions require independent professional review','Formal cover layout and signature area remain pending formal template confirmation'],'family':data['family'],'source_sha256':profile['source_sha256'],'format_policy_sha256':sha(POLICY_PATH),'format_policy_status':policy['status'],'company_identity':{'name':policy['company_identity']['name'],'logo_sha256':policy['company_identity']['logo_sha256']},'toc_depth':max(heading_levels,default=1),'output_sha256':sha(out),'block_ids':[b['id'] for b in data['blocks']],'bookmark_map':bookmark_map}
    from handoff import encoded
    report['content_sha256']=hashlib.sha256(encoded(data)).hexdigest()
    report['content_hash_method']='sha256-normalized-content-v1'
    jwrite(Path(out).with_suffix('.quality.json'),report); return report

def prose(b):
    t=b['type']
    if t in ('heading','paragraph'): return b['text']
    if t=='runs': return ''.join(r['text'] for r in b['runs'])
    if t=='list': return '\n'.join(b['items'])
    if t=='table': return b['caption']+'\n'+' | '.join(b['columns'])+'\n'+'\n'.join(' | '.join(r) for r in b['rows'])
    if t=='form': return b['caption']+'\n'+'\n'.join(f'[{c["row"]},{c["cell"]}] '+ '\n'.join(c['paragraphs']) for c in b['cells'])
    if t=='figure': return b['caption']+' ['+b['alt']+']'
    return ''
def markdown(data,out):
    check_new_outputs(out)
    lines=['# '+data['metadata']['title'],data['metadata']['subtitle'],'状态：'+data['metadata']['status']]
    for b in data['blocks']:
        lines+=['',f'<!-- {b["id"]}; sources={",".join(b["source_ids"])}; requirements={",".join(b["requirement_ids"])} -->',('#'*(b['level']+1)+' ' if b['type']=='heading' else '')+prose(b)]
    lines+=['','## 待确认事项']+[p['id']+' '+p['text'] for p in data['pending']]
    Path(out).write_text('\n'.join(lines),encoding='utf-8')
def html_demo(data,out):
    check_new_outputs(out)
    css=(ROOT/'assets/theme.css').read_text(encoding='utf-8'); esc=html.escape
    blocks=[]
    for b in data['blocks']:
        ident=esc(b['id']); typ=b['type']; s=''
        if typ=='heading': s=f'<h{b["level"]+1}>{esc(b["text"])}</h{b["level"]+1}>'
        elif typ=='table': s='<p>'+esc(b['caption'])+'</p><table><thead><tr>'+''.join('<th>'+esc(v)+'</th>' for v in b['columns'])+'</tr></thead><tbody>'+''.join('<tr>'+''.join('<td>'+esc(v)+'</td>' for v in r)+'</tr>' for r in b['rows'])+'</tbody></table>'
        else: s='<p>'+esc(prose(b)).replace('\n','<br>')+'</p>'
        blocks.append(f'<section id="{ident}" data-sources="{esc(",".join(b["source_ids"]))}">{s}</section>')
    Path(out).write_text('<!doctype html><html lang="zh-CN"><meta charset="utf-8"><meta name="viewport" content="width=device-width"><meta http-equiv="Content-Security-Policy" content="default-src \'none\'; style-src \'unsafe-inline\'"><title>'+esc(data['metadata']['title'])+'</title><style>'+css+'</style><main><h1>'+esc(data['metadata']['title'])+'</h1><p>'+esc(data['metadata']['subtitle'])+'</p>'+''.join(blocks)+'<h2>待确认事项</h2>'+''.join('<p>'+esc(p['id']+' '+p['text'])+'</p>' for p in data['pending'])+'</main></html>',encoding='utf-8')
def slides(data,out):
    from presentations import render
    render(data,out)

def main():
    ap=argparse.ArgumentParser(); ap.add_argument('format',choices=['docx','md','html','pptx','validate']); ap.add_argument('content'); ap.add_argument('output',nargs='?')
    a=ap.parse_args(); data=read_content(a.content)
    if a.format=='validate': print('VALID'); return
    if not a.output: ap.error('output required')
    out=Path(a.output).resolve(); out.parent.mkdir(parents=True,exist_ok=True)
    if out==Path(a.content).resolve() or out.is_relative_to(ROOT/'assets'): raise ValueError('refusing to overwrite input or assets')
    if a.format=='docx': word(data,out,Path(a.content).resolve().parent)
    elif a.format=='md': markdown(data,out)
    elif a.format=='html': html_demo(data,out)
    else: slides(data,out)
if __name__=='__main__': main()
